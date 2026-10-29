# -*- coding: utf-8 -*-
"""
Checks the recalculated workbook against reference_calc.py.

Run after Excel has recalculated merit-model.xlsx (recalc_model.ps1 does both).
Exits 1 on any formula error or any mismatch above one kopeck.
"""
import sys

from openpyxl import load_workbook

from reference_calc import MODEL, SAMPLE_EMPLOYEES, SAMPLE_RANGES, analyse, load, load_bands

SRC = sys.argv[1] if len(sys.argv) > 1 else SAMPLE_EMPLOYEES
BANDS = sys.argv[2] if len(sys.argv) > 2 else SAMPLE_RANGES
XLSX = sys.argv[3] if len(sys.argv) > 3 else MODEL

sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(XLSX, data_only=True)
fail = 0

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#"):
                print("ERROR %s!%s %s" % (ws.title, c.coordinate, c.value))
                fail += 1

ref = analyse(load(SRC), load_bands(BANDS))

# the summary sheet carries its own cell map: key in column D, address in column E
s = next((ws for ws in wb.worksheets if ws["D1"].value == "cells_for_check"), None)
if s is None:
    print("FAIL no sheet with the cells_for_check marker in D1")
    sys.exit(1)
cell = {}
for r in range(2, s.max_row + 1):
    k, a = s.cell(r, 4).value, s.cell(r, 5).value
    if k and a:
        cell[k] = a

expect = {
    "merit_base": ref["merit"]["base"], "merit_bonus": ref["merit"]["bonus"],
    "merit_contrib": ref["merit"]["contrib"], "merit_annual": ref["merit"]["annual"],
    "merit_in_year": ref["merit"]["in_year"], "target_annual": ref["target"]["annual"],
    "target_in_year": ref["target"]["in_year"], "pay_before": ref["pay_before"],
    "contrib_before": ref["contrib_before"], "payroll": ref["payroll"],
    "eligible_base": ref["eligible_base"], "merit_pct": ref["merit_pct"] / 100,
    "dev_rub": ref["dev_rub"], "dev_rub_in_year": ref["dev_rub_in_year"],
    "below_min_after": ref["below_min_after"], "red_circled": ref["red_circled"],
    "eligible": sum(1 for e in ref["emp"] if e["eligible"]),
}
for key, want in expect.items():
    got = s[cell[key]].value if key in cell else None
    tol = 1e-6 if abs(want) < 1 else 0.01  # shares vs rouble amounts
    ok = isinstance(got, (int, float)) and abs(got - want) < tol
    print("%-4s %-16s excel=%18.4f ref=%18.4f" % ("OK" if ok else "FAIL", key, got if ok or isinstance(got, (int, float)) else 0, want))
    fail += 0 if ok else 1

c = next(ws for ws in wb.worksheets if ws.title == "Расчёт")
hdr = {c.cell(3, j).value: j for j in range(1, c.max_column + 1)}
for i, e in enumerate(ref["emp"]):
    r = 4 + i
    for h, key in (("Повышение", "increase"), ("Зона CR", "zone"), ("Участвует", "eligible")):
        got, want = c.cell(r, hdr[h]).value, float(e[key])
        if not isinstance(got, (int, float)) or abs(got - want) > 0.01:
            print("FAIL %s %s excel=%s ref=%s" % (e["id"], h, got, want))
            fail += 1
    monthly = (("Оклад в месяц до пересмотра", e["base"] / 12), ("Оклад в месяц после пересмотра", e["new_base"] / 12),
               ("Переменная часть в месяц до пересмотра", e["base"] * e["bonus"] / 12),
               ("Переменная часть в месяц после пересмотра", e["new_base"] * e["bonus"] / 12),
               ("Совокупный доход в месяц до пересмотра", e["base"] * (1 + e["bonus"]) / 12),
               ("Совокупный доход в месяц после пересмотра", e["new_base"] * (1 + e["bonus"]) / 12))
    for h, want in monthly:
        got = c.cell(r, hdr[h]).value
        if not isinstance(got, (int, float)) or abs(got - want) > 0.01:
            print("FAIL %s %s excel=%s ref=%s" % (e["id"], h, got, want))
            fail += 1

print("employees checked:", len(ref["emp"]))
print("FAILURES:", fail)
sys.exit(1 if fail else 0)
