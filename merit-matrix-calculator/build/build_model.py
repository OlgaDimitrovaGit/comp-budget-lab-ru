# -*- coding: utf-8 -*-
"""
Builds merit-model.xlsx from a CSV of employees.

Every figure is a live formula. The calculation sheet runs left to right in
numbered steps, one simple operation per column; rows 1-3 carry the step, a
note and the column header. Sheet names and labels are Russian: the workbook
is for a Russian-speaking reader.
"""
import csv
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from reference_calc import DEFAULTS, MODEL, SAMPLE_EMPLOYEES, SAMPLE_RANGES, load_bands, parse_date

SRC = sys.argv[1] if len(sys.argv) > 1 else SAMPLE_EMPLOYEES
BANDS = sys.argv[2] if len(sys.argv) > 2 else SAMPLE_RANGES
OUT = sys.argv[3] if len(sys.argv) > 3 else MODEL

INPUT = PatternFill("solid", fgColor="FFF4CC")
AMBER = PatternFill("solid", start_color="FFC000", end_color="FFC000")
STEP = PatternFill("solid", fgColor="DCE6F0")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
RUB = '#,##0.00" ₽"'
PCT = "0.00%"
DATE = "dd.mm.yyyy"
MONTH = "mmmm yyyy"
MONTHS = '0 "мес."'

# sheet names; quoted in formulas because they are Cyrillic
SH_IN, SH_LISTS, SH_BANDS, SH_EMP, SH_CALC, SH_MATRIX, SH_SUM = (
    "Настройки", "Списки", "Вилки", "Сотрудники", "Расчёт", "Матрица", "Итоги")


def q(sheet):
    return f"'{sheet}'!"


YES, NO, NO_CAP = "Да", "Нет", "Без ограничения"

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


# ---------------------------------------------------------------- Inputs
ws = wb.active
ws.title = SH_IN
ws["A1"] = "Настройки: меняйте жёлтые ячейки"
ws["A1"].font = Font(bold=True, size=13)
ws.append([])
ws.append(["Параметр", "Значение", "Пояснение"])
for c in ws[3]:
    c.font = BOLD

S = DEFAULTS
settings = [
    ("EffDate", "Дата вступления в силу (выберите месяц)", S["effective_date"], MONTH, True,
     "Повышения действуют с 1-го числа этого месяца. Право на пересмотр считается на эту дату. Выберите из списка."),
    ("EffMonth", "Месяц вступления в силу", "=MONTH(EffDate)", "0", False,
     "Номер месяца даты вступления в силу."),
    ("InYearFactor", "Доля года", "=(13-EffMonth)/12", PCT, False,
     "Месяцы с месяца вступления в силу по декабрь включительно, делённые на 12: повышение начинается 1-го числа."),
    ("PoolPct", "Целевой мерит, % от ФОТ окладов участников", S["pool_pct"] / 100, PCT, True,
     "Плановое среднее повышение в процентах от годовых окладов участников пересмотра, с учётом ставки (FTE). "
     "Целевой бюджет на листе «Итоги» — этот процент каждому участнику, вместе с переменной частью и взносами."),
    ("MinMonthsHire", "Минимальный стаж, мес.", S["min_months_hire"], MONTHS, True,
     "Принят меньше чем за столько месяцев до даты вступления в силу — в пересмотре не участвует. "
     "Принят после этой даты — не участвует."),
    ("MinMonthsChange", "Без индивидуального изменения оклада за последние, мес.", S["min_months_change"], MONTHS, True,
     "Оклад меняли индивидуально меньше чем за столько месяцев до даты вступления в силу — в пересмотре не участвует."),
    ("MaxIncrease", "Потолок повышения на человека, % от оклада", S["max_increase"] or NO_CAP, PCT, True,
     "Страховочный предел. Повышение — разрыв до целевого CR или минимум рейтинга, если он больше; "
     "потолок срезает его, только если оно выше. «Без ограничения» (или пустая ячейка, или 0%) — предела нет."),
    ("ClipAtMax", "Не поднимать выше максимума вилки", YES if S["clip_at_max"] else NO, "@", True,
     "Да: повышение никогда не выводит оклад выше максимума вилки."),
    ("CompanyTarget", "Политика оплаты, % от середины вилки", S["company_target"], PCT, True,
     "Оставьте 100%, если середины вилок уже отражают вашу политику оплаты."),
    ("RateBelow", "Ставка страховых взносов до предельной базы", S["rate_below"] / 100, PCT, True,
     "Россия: единый тариф 30% до предельной базы (п. 3 ст. 425 НК РФ). "
     "Взносы на травматизм не включены: их ставка зависит от класса риска, добавьте её сюда при необходимости. "
     "Предельная база на 2027 год ещё не утверждена: взята база 2026 года."),
    ("RateAbove", "Ставка страховых взносов сверх предельной базы", S["rate_above"] / 100, PCT, True,
     "Россия: 15,1% с выплат сверх предельной базы (п. 3 ст. 425 НК РФ)."),
    ("Ceiling", "Предельная база в год на человека", S["ceiling"], RUB, True,
     "Россия, 2026: 2 979 000 ₽ (Постановление Правительства РФ от 31.10.2025 № 1705)."),
]
for i, (n, label, val, fmt, is_input, note) in enumerate(settings):
    r = 4 + i
    ws.cell(r, 1, label)
    c = ws.cell(r, 2, val)
    c.number_format = fmt
    if is_input:
        c.fill = INPUT
    ws.cell(r, 3, note).alignment = WRAP
    name(n, f"{q(SH_IN)}$B${r}")

RT = 4 + len(settings) + 2  # rating table header row
ws.cell(RT - 1, 1, "Цели по рейтингу результативности (условное допущение: задайте свои)").font = BOLD
for j, h in enumerate(["Рейтинг", "Название", "Целевой CR", "Минимальное повышение, % от оклада"]):
    ws.cell(RT, 1 + j, h).font = BOLD
labels = {1: "Не соответствует", 2: "Частично соответствует", 3: "Соответствует", 4: "Превосходит",
          5: "Выдающийся"}
for k in range(1, 6):
    r = RT + k
    tgt, mn = S["ratings"][k]
    ws.cell(r, 1, k)
    ws.cell(r, 2, labels[k])
    for j, v in enumerate((tgt, mn)):
        c = ws.cell(r, 3 + j, v)
        c.number_format = PCT
        c.fill = INPUT
R1, R2 = RT + 1, RT + 5
notes = [
    "Целевой CR (compa-ratio, оклад к середине вилки): до него поднимают оклад при этом рейтинге. Повышение — "
    "разрыв от оклада до середины × политика оплаты × целевой CR. 0% — рейтинг без повышения. Цель выше "
    "максимума вилки (максимум / середина) срезается по максимуму, когда «Не поднимать выше максимума вилки» = Да.",
    "Минимальное повышение: нижний предел для рейтинга, платится даже тем, кто уже на цели или выше. 0% — без "
    "предела. Не действует на максимуме вилки и выше и для тех, кто не участвует в пересмотре; ограничен "
    "потолком на человека и, когда включено «не выше максимума», запасом до максимума.",
]
for i, note in enumerate(notes):
    ws.cell(R2 + 1 + i, 1, note).alignment = WRAP

ZT = R2 + 5  # zone table header row
ws.cell(ZT - 1, 1, "Зоны CR: только для вида на листе «Матрица»").font = BOLD
ws.cell(ZT - 1, 3, "Повышения от зон не зависят. Зоны лишь раскладывают людей в классическую сетку "
        "«рейтинг × CR» на листе «Матрица».").alignment = WRAP
for j, h in enumerate(["Зона", "Нижняя граница (CR от)"]):
    ws.cell(ZT, 1 + j, h).font = BOLD
for k, z in enumerate(S["zones"]):
    ws.cell(ZT + 1 + k, 1, k + 1)
    c = ws.cell(ZT + 1 + k, 2, z)
    c.number_format = PCT
    c.fill = INPUT
Z1, Z2 = ZT + 1, ZT + len(S["zones"])
for col, w in zip("ABC", (50, 18, 80)):
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- Lists (drop-down values for Inputs)
wl = wb.create_sheet(SH_LISTS)
LISTS = {
    "EffDate": ("Дата вступления в силу", [date(2026 + (10 + i) // 12, (10 + i) % 12 + 1, 1) for i in range(26)], MONTH),
    "MinMonthsHire": ("Минимальный стаж", [0, 1, 2, 3, 6, 9, 12], MONTHS),
    "MinMonthsChange": ("Без изменения оклада", [0, 1, 2, 3, 6, 9, 12], MONTHS),
    "MaxIncrease": ("Потолок повышения", [NO_CAP, 0.10, 0.15, 0.20, 0.25, 0.30, 0.40, 0.50, 0.75, 1.00], PCT),
    "ClipAtMax": ("Не выше максимума", [YES, NO], "@"),
}
setting_row = {n: 4 + i for i, (n, *_) in enumerate(settings)}
for j, (key, (head, values, fmt)) in enumerate(LISTS.items()):
    col = get_column_letter(1 + j)
    wl.cell(1, 1 + j, head).font = BOLD
    for i, v in enumerate(values):
        wl.cell(2 + i, 1 + j, v).number_format = fmt
    wl.column_dimensions[col].width = 20
    dv = DataValidation(type="list", formula1=f"={q(SH_LISTS)}${col}$2:${col}${1 + len(values)}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Выберите из списка",
                        error="Выберите значение из раскрывающегося списка.")
    ws.add_data_validation(dv)
    dv.add(f"B{setting_row[key]}")
assert S["effective_date"] in LISTS["EffDate"][1]
wl.sheet_state = "hidden"

# ---------------------------------------------------------------- Bands
wb_b = wb.create_sheet(SH_BANDS)
wb_b.append(["Категория", "Грейд", "Минимум вилки", "Середина вилки", "Максимум вилки", "Ключ поиска"])
for c in wb_b[1]:
    c.font = BOLD
for i, ((cat, g), (mn, md, mx)) in enumerate(load_bands(BANDS).items()):
    r = 2 + i
    wb_b.append([cat, g, mn, md, mx, f'=A{r}&"|"&B{r}'])
CAP_BANDS = max(500, wb_b.max_row - 1)
for r in range(wb_b.max_row + 1, CAP_BANDS + 2):
    wb_b.cell(r, 6, f'=IF(A{r}="","",A{r}&"|"&B{r})')
for row in wb_b.iter_rows(min_row=2, max_row=CAP_BANDS + 1, max_col=5):
    for c in row:
        c.fill = INPUT
    for c in row[2:]:
        c.number_format = RUB
wb_b.column_dimensions["A"].width = 22
for col in "BCDEF":
    wb_b.column_dimensions[col].width = 16

# ---------------------------------------------------------------- Employees
we = wb.create_sheet(SH_EMP)
heads = ["ID", "Категория", "Грейд", "Рейтинг", "Дата приёма", "Дата последнего изменения оклада",
         "Годовой оклад (полная ставка)", "Переменная часть, % от оклада", "Ставка (FTE)"]
we.append(heads)
for c in we[1]:
    c.font = BOLD
    c.alignment = WRAP
rows = list(csv.DictReader(open(SRC, encoding="utf-8-sig")))
for r in rows:
    we.append([r["id"], r["category"], int(r["grade"]), int(r["rating"]),
               parse_date(r["hire_date"]), parse_date(r["last_change_date"]),
               float(r["base_salary"]), float(r["bonus_pct"]), float(r["fte"])])
N = len(rows)
# Room for your own staff: Calc and Bands carry formulas for more rows than the demo has.
# A row with no ID stays blank and adds nothing.
CAP_EMP = max(2000, N)
for row in we.iter_rows(min_row=2, max_row=CAP_EMP + 1, max_col=len(heads)):
    for c in row:
        c.fill = INPUT
    row[4].number_format = row[5].number_format = DATE
    row[6].number_format = RUB
for col, w in zip("ABCDEFGHI", (8, 22, 8, 8, 12, 16, 16, 14, 9)):
    we.column_dimensions[col].width = w
we.row_dimensions[1].height = 45

# ---------------------------------------------------------------- Calc
def EMP(field):
    """Look a field up on Employees by ID and by column header, not by cell."""
    e = q(SH_EMP)
    return f'=INDEX({e}$A:$Z,MATCH({{id}},{e}$A:$A,0),MATCH("{field}",{e}$1:$1,0))'


# Employer contributions on annual pay {x}: one rate up to the ceiling, another above it
CONTRIB = "=RateBelow*MIN({x},Ceiling)+RateAbove*MAX(0,{x}-Ceiling)"

B = q(SH_BANDS)
I = q(SH_IN)
BAND_KEY = '{cat}&"|"&{grade}'

# (key, header, number format, formula template, note); {k} -> this row's cell of column k
CALC = [
    ("Шаг 0. Исходные данные", [
        ("n", "№", "0", "=ROW()-3", "Номер строки в списке."),
        ("id", "ID", None, f'=IF(INDEX({q(SH_EMP)}$A:$A,{{n}}+1)="","",INDEX({q(SH_EMP)}$A:$A,{{n}}+1))',
         "n-й ID с листа «Сотрудники». Пусто, если список короче: тогда пуста вся строка."),
        ("cat", "Категория", None, EMP("Категория"), "По ID с листа «Сотрудники»."),
        ("grade", "Грейд", "0", EMP("Грейд"), ""),
        ("rating", "Рейтинг", "0", EMP("Рейтинг"), "От 1 до 5."),
        ("hire", "Дата приёма", DATE, EMP("Дата приёма"), ""),
        ("chg", "Последнее изменение оклада", DATE, EMP("Дата последнего изменения оклада"),
         "Последнее индивидуальное изменение."),
        ("base", "Годовой оклад", RUB, EMP("Годовой оклад (полная ставка)"), "Годовой оклад до вычета НДФЛ, на полную ставку."),
        ("bonus", "Переменная часть, % от оклада", PCT, EMP("Переменная часть, % от оклада") + "/100",
         "Доля переменной части от оклада."),
        ("fte", "Ставка", "0.00", EMP("Ставка (FTE)"), "1 — полная ставка."),
    ]),
    ("Шаг 1. Вилка и годовой доход", [
        ("bok", "Вилка найдена", "0", f'=IF(ISNA(MATCH({BAND_KEY},{B}$F:$F,0)),0,1)',
         "0: на листе «Вилки» нет строки для этой категории и грейда; итоги покажут #Н/Д, пока её не добавят."),
        ("bmin", "Минимум вилки", RUB, f"=INDEX({B}$C:$C,MATCH({BAND_KEY},{B}$F:$F,0))",
         "С листа «Вилки» по категории и грейду."),
        ("bmid", "Середина вилки", RUB, f"=INDEX({B}$D:$D,MATCH({BAND_KEY},{B}$F:$F,0))", ""),
        ("bmax", "Максимум вилки", RUB, f"=INDEX({B}$E:$E,MATCH({BAND_KEY},{B}$F:$F,0))", ""),
        ("tot", "Оклад + переменная часть", RUB, "={base}*(1+{bonus})", "На полную ставку."),
    ]),
    ("Шаг 2. Право на пересмотр", [
        ("mh", "Месяцев с приёма", "0", '=IF({hire}>EffDate,-1,DATEDIF({hire},EffDate,"m"))',
         "Полные месяцы до даты вступления в силу; −1, если принят после неё."),
        ("mc", "Месяцев с изменения оклада", "0", '=IF({chg}>EffDate,-1,DATEDIF({chg},EffDate,"m"))',
         "−1, если изменение позже даты вступления в силу."),
        ("elig", "Участвует", "0", "=IF(AND({mh}>=MinMonthsHire,{mc}>=MinMonthsChange),1,0)", "1 — участвует в пересмотре."),
    ]),
    ("Шаг 3. Целевой оклад", [
        ("cr", "CR", PCT, "={base}/{bmid}", "Compa-ratio: оклад / середина вилки."),
        ("zone", "Зона CR", "0", f"=MATCH({{cr}},{I}$B${Z1}:$B${Z2},1)",
         "В расчёт повышения не входит. Только ставит человека в столбец сетки на листе «Матрица» (рейтинг × CR)."),
        ("tcr", "Целевой CR для рейтинга", PCT,
         f"=INDEX({I}$C${R1}:$C${R2},MATCH({{rating}},{I}$A${R1}:$A${R2},0))", "Из таблицы рейтингов."),
        ("minp", "Минимальное повышение", PCT,
         f"=INDEX({I}$D${R1}:$D${R2},MATCH({{rating}},{I}$A${R1}:$A${R2},0))",
         "Нижний предел для рейтинга, платится и на цели, и выше; см. «Настройки»."),
        ("ts", "Целевой оклад", RUB, "={bmid}*CompanyTarget*{tcr}", "Середина × политика оплаты × целевой CR."),
        ("gap", "Разрыв до цели", RUB, "=MAX(0,{ts}-{base})", "0, если оклад уже на цели или выше."),
    ]),
    ("Шаг 4. Мерит", [
        ("capon", "Потолок действует", "0", "=IF(AND(ISNUMBER(MaxIncrease),MaxIncrease>0),1,0)",
         "0: «Без ограничения», пусто или 0% на листе «Настройки»."),
        ("capa", "Потолок повышения, ₽", RUB, "=IF({capon}=1,{base}*MaxIncrease,0)",
         "Наибольшее повышение для человека: потолок × оклад; 0, если потолка нет."),
        ("mina", "Минимальное повышение, ₽", RUB, "={base}*{minp}", ""),
        ("rule", "Повышение до ограничений вилки", RUB,
         "=IF({tcr}=0,0,IF({capon}=1,MIN(MAX({gap},{mina}),{capa}),MAX({gap},{mina})))",
         "Разрыв до цели, не меньше минимума и не больше потолка (если он есть)."),
        ("red", "На максимуме вилки или выше", "0", "=IF({base}>={bmax},1,0)",
         "На максимуме вилки или выше: мерита нет."),
        ("afr", "После правила максимума", RUB, "=IF({red}=1,0,{rule})",
         "Оклад уже на максимуме или выше — повышение 0, что бы ни дало правило."),
        ("room", "Запас до максимума", RUB, "=MAX(0,{bmax}-{base})",
         "Максимум вилки − оклад: сколько можно добавить, не выходя из вилки."),
        ("afc", "После среза по максимуму", RUB, f'=IF(ClipAtMax="{NO}",{{afr}},MIN({{afr}},{{room}}))',
         "Если «Не поднимать выше максимума вилки» = Да, повышение больше запаса срезается до запаса: новый оклад = максимум."),
        ("inc", "Мерит", RUB, "=IF({elig}=1,{afc},0)", "Прибавка к годовому окладу на полную ставку."),
        ("incp", "Мерит, %", PCT, "=IF({base}>0,{inc}/{base},0)", ""),
        ("nb", "Новый оклад", RUB, "={base}+{inc}", ""),
        ("bmn", "Ниже минимума вилки после пересмотра", "0", "=IF({nb}<{bmin},1,0)", "Флаг."),
    ]),
    ("Шаг 5. Стоимость (в год, с учётом ставки)", [
        ("p0", "Оклад + переменная часть до пересмотра, по ставке", RUB, "={tot}*{fte}",
         "(Годовой оклад + переменная часть) × ставка."),
        ("c0", "Страховые взносы до пересмотра", RUB, CONTRIB.replace("{x}", "{p0}"),
         "Одна ставка до предельной базы, другая сверх неё (лист «Настройки»)."),
        ("p2", "Оклад + переменная часть после мерита, по ставке", RUB, "={nb}*(1+{bonus})*{fte}", ""),
        ("c2", "Страховые взносы после мерита", RUB, CONTRIB.replace("{x}", "{p2}"), ""),
        ("s2b", "Мерит: оклад", RUB, "={inc}*{fte}", ""),
        ("s2n", "Мерит: переменная часть", RUB, "={inc}*{bonus}*{fte}", "Переменная часть растёт вместе с окладом."),
        ("s2c", "Мерит: страховые взносы", RUB, "={c2}-{c0}", "Дополнительные взносы с прибавки."),
        ("eb", "Оклад участника", RUB, "=IF({elig}=1,{base}*{fte},0)",
         "Оклад участника пересмотра с учётом ставки: база для целевого процента."),
    ]),
    ("Шаг 5а. Целевой бюджет: целевой процент каждому участнику", [
        ("tinc", "Целевое повышение", RUB, "=IF({elig}=1,{base}*PoolPct,0)",
         "Оклад × целевой процент, на полную ставку. Ориентир, а не выплата."),
        ("s3b", "Цель: оклад", RUB, "={tinc}*{fte}", ""),
        ("s3n", "Цель: переменная часть", RUB, "={tinc}*{bonus}*{fte}", "У каждого своя доля переменной части."),
        ("pT", "Оклад + переменная часть по цели, по ставке", RUB, "=({base}+{tinc})*(1+{bonus})*{fte}", ""),
        ("cT", "Страховые взносы по цели", RUB, CONTRIB.replace("{x}", "{pT}"), ""),
        ("s3c", "Цель: страховые взносы", RUB, "={cT}-{c0}",
         "Зависит от того, где человек относительно предельной базы, поэтому это не одна ставка от целевого оклада."),
    ]),
]

wc = wb.create_sheet(SH_CALC)
col_of = {}
ci = 1
for step, items in CALC:
    for key, *_ in items:
        col_of[key] = get_column_letter(ci)
        ci += 1
DATA0 = 4

ci = 1
for step, items in CALC:
    first = True
    for key, header, fmt, tmpl, note in items:
        if first:
            wc.cell(1, ci, step).font = BOLD
            first = False
        wc.cell(1, ci).fill = STEP
        wc.cell(2, ci, note).alignment = WRAP
        wc.cell(3, ci, header).font = BOLD
        wc.cell(3, ci).alignment = WRAP
        if key not in ("n", "id"):
            tmpl = '=IF({id}="","",' + tmpl[1:] + ")"
        for i in range(CAP_EMP):
            r = DATA0 + i
            c = wc.cell(r, ci, tmpl.format(**{k: f"{v}{r}" for k, v in col_of.items()}))
            if fmt:
                c.number_format = fmt
        wc.column_dimensions[get_column_letter(ci)].width = 14
        ci += 1
wc.row_dimensions[2].height = 90
wc.row_dimensions[3].height = 45
wc.freeze_panes = "B4"


def CC(key):
    """Whole Calc column, for SUMIFS/COUNTIFS."""
    L = col_of[key]
    return f"{q(SH_CALC)}${L}:${L}"


# ---------------------------------------------------------------- Matrix
wm = wb.create_sheet(SH_MATRIX)
wm["A1"] = "Матрица мерита: результат, посчитанный по каждому человеку"
wm["A1"].font = Font(bold=True, size=13)
wm["A2"] = "Строки — рейтинг, столбцы — зона CR (оклад / середина вилки). Только участники пересмотра."
blocks = [("Мерит, % от окладов участников", PCT, "pct"), ("Численность", "0", "n"),
          ("Мерит, ₽ (в год, с учётом ставки)", RUB, "rub")]
r0 = 4
nz = len(S["zones"])
for title, fmt, kind in blocks:
    wm.cell(r0, 1, title).font = BOLD
    for z in range(1, nz + 1):
        lo = f'TEXT(INDEX({I}$B${Z1}:$B${Z2},{z}),"0%")'
        hi = f'TEXT(INDEX({I}$B${Z1}:$B${Z2},{z + 1}),"0%")'
        head = f'="CR < "&{hi}' if z == 1 else f'="CR ≥ "&{lo}' if z == nz else f'="CR от "&{lo}&" до "&{hi}'
        wm.cell(r0 + 1, 1 + z, head).font = BOLD
    for k in range(1, 6):
        r = r0 + 1 + k
        wm.cell(r, 1, f'=INDEX({I}$A${R1}:$A${R2},{k})&" "&INDEX({I}$B${R1}:$B${R2},{k})')
        rk = f"INDEX({I}$A${R1}:$A${R2},{k})"
        for z in range(1, nz + 1):
            crit = f'{CC("rating")},{rk},{CC("zone")},{z},{CC("elig")},1'
            if kind == "n":
                f = f"=COUNTIFS({crit})"
            elif kind == "rub":
                f = f"=SUMIFS({CC('s2b')},{crit})"
            else:
                f = f"=IF(SUMIFS({CC('eb')},{crit})>0,SUMIFS({CC('s2b')},{crit})/SUMIFS({CC('eb')},{crit}),0)"
            wm.cell(r, 1 + z, f).number_format = fmt
    r0 += 9
wm.column_dimensions["A"].width = 30
for col in "BCDE":
    wm.column_dimensions[col].width = 18

# ---------------------------------------------------------------- Summary
wsum = wb.create_sheet(SH_SUM, 0)
wsum["A1"] = "Итоги: во что обходится пересмотр"
wsum["A1"].font = Font(bold=True, size=13)
wsum.append([])
wsum.append(["", "Значение"])
for c in wsum[3]:
    c.font = BOLD


def S_(key):
    return f"SUM({CC(key)})"


# (key, label, formula, number format); {key} -> that line's value cell
lines = [
    ("hc", "Численность", f"=COUNT({CC('base')})", "0"),
    ("ne", "Участвуют в пересмотре", f"={S_('elig')}", "0"),
    (None,),
    (None, "РАСХОДЫ НА ОПЛАТУ ТРУДА СО СТРАХОВЫМИ ВЗНОСАМИ, полный год", None, None),
    ("pr0", "Расходы сейчас", f"={S_('p0')}+{S_('c0')}", RUB),
    ("pay0", "   в т.ч. ФОТ: оклад + переменная часть", f"={S_('p0')}", RUB),
    ("con0", "   в т.ч. страховые взносы", f"={S_('c0')}", RUB),
    ("pr2", "Расходы после пересмотра", f"={S_('p2')}+{S_('c2')}", RUB),
    ("d2", "Прирост за полный год", "={pr2}-{pr0}", RUB),
    ("dtp", "Прирост, % от расходов сейчас", "=IF({pr0}>0,{d2}/{pr0},0)", PCT),
    (None,),
    (None, "В ЭТОМ ГОДУ: с месяца вступления в силу", None, None),
    ("tt", "Прирост в этом году", "={d2}*InYearFactor", RUB),
    ("ttp", "Прирост в этом году, % от расходов сейчас", "=IF({pr0}>0,{tt}/{pr0},0)", PCT),
    (None,),
    (None, "МЕРИТ (лист «Расчёт», шаг 4): из чего складывается", None, None),
    ("b_b", "Мерит: прибавка к окладам", f"={S_('s2b')}", RUB),
    ("b_n", "Рост переменной части", f"={S_('s2n')}", RUB),
    ("b_c", "Страховые взносы", f"={S_('s2c')}", RUB),
    ("b_y", "Итого за полный год", "={b_b}+{b_n}+{b_c}", RUB),
    ("b_t", "Итого в этом году", "={b_y}*InYearFactor", RUB),
    (None,),
    (None, "МЕРИТ ПРОТИВ ЦЕЛЕВОГО БЮДЖЕТА, с переменной частью и взносами", None, None),
    ("eb", "Оклады участников пересмотра", f"={S_('eb')}", RUB),
    ("pp", "Целевой мерит, % от окладов участников (меняется на листе «Настройки»)", "=PoolPct", PCT),
    ("mp", "Мерит по расчёту, % от окладов участников", "=IF({eb}>0,{b_b}/{eb},0)", PCT),
    ("dpp", "Отклонение от цели, п.п. (+ перерасход, − экономия)", "=({mp}-{pp})*100", "0.00"),
    ("g_y", "Целевой бюджет за полный год", "={g_b}+{g_n}+{g_c}", RUB),
    ("g_b", "   в т.ч. оклады", f"={S_('s3b')}", RUB),
    ("g_n", "   в т.ч. переменная часть", f"={S_('s3n')}", RUB),
    ("g_c", "   в т.ч. страховые взносы", f"={S_('s3c')}", RUB),
    ("deur", "Отклонение от целевого бюджета, ₽ за полный год (+ перерасход)", "={b_y}-{g_y}", RUB),
    ("g_t", "Целевой бюджет в этом году", "={g_y}*InYearFactor", RUB),
    ("deut", "Отклонение от целевого бюджета, ₽ в этом году (+ перерасход)", "={b_t}-{g_t}", RUB),
    (None,),
    (None, "ФЛАГИ", None, None),
    ("f2", "На максимуме вилки или выше: мерита нет", f"={S_('red')}", "0"),
    ("f3", "Ниже минимума вилки после пересмотра", f"={S_('bmn')}", "0"),
    (None,),
    (None, "ПРОВЕРКИ ДАННЫХ: везде должен быть 0", None, None),
    ("k1", "Строки «Сотрудников», не попавшие в расчёт", f"=COUNTA({q(SH_EMP)}$A:$A)-1-{{hc}}", "0"),
    ("k2", "Сотрудники без вилки на листе «Вилки»", f"={{hc}}-{S_('bok')}", "0"),
    ("k4", "Сотрудники с рейтингом, которого нет в таблице рейтингов",
     f"={{hc}}-SUMPRODUCT(COUNTIFS({CC('rating')},{I}$A${R1}:$A${R2}))", "0"),
]
srow = {ln[0]: 4 + i for i, ln in enumerate(lines) if len(ln) > 1 and ln[0]}
for i, ln in enumerate(lines):
    r = 4 + i
    if len(ln) == 1:
        continue
    key, label, f, fmt = ln
    wsum.cell(r, 1, label)
    if f is None:
        wsum.cell(r, 1).font = BOLD
        continue
    c = wsum.cell(r, 2, f.format(**{k: f"B{v}" for k, v in srow.items()}))
    c.number_format = fmt
wsum.column_dimensions["A"].width = 66
wsum.column_dimensions["B"].width = 20
c = wsum.cell(srow["pp"], 1)
c.hyperlink = f"#{q(SH_IN)}B{setting_row['PoolPct']}"
c.font = Font(color="0563C1", underline="single")
for k in ("k1", "k2", "k4"):
    wsum.conditional_formatting.add(f"B{srow[k]}", CellIsRule(operator="notEqual", formula=["0"], fill=AMBER))

# Cell map for check_model.py: key -> Summary cell
CHECK = {"merit_base": "b_b", "merit_bonus": "b_n", "merit_contrib": "b_c", "merit_annual": "b_y",
         "merit_in_year": "b_t", "target_annual": "g_y", "target_in_year": "g_t", "pay_before": "pay0",
         "contrib_before": "con0", "payroll": "pr0", "eligible_base": "eb", "merit_pct": "mp",
         "dev_rub": "deur", "dev_rub_in_year": "deut", "below_min_after": "f3", "red_circled": "f2",
         "eligible": "ne"}
wsum["D1"] = "cells_for_check"
wsum["D1"].font = Font(color="A6A6A6")
for i, (k, s) in enumerate(CHECK.items()):
    wsum.cell(2 + i, 4, k).font = Font(color="A6A6A6")
    wsum.cell(2 + i, 5, f"B{srow[s]}").font = Font(color="A6A6A6")

wb.save(OUT)
print("written", OUT, "rows", N)
