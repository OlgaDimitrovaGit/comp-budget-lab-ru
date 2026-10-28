# -*- coding: utf-8 -*-
"""
Builds the Excel counterpart of the plan-vs-actual dashboard (Russian edition).

Reads budget-plan-fact.csv and department-deviation.csv and produces an .xlsx
in which the whole chain is live: plan -> actual -> deviation -> running total
-> forecast on each department's own trend -> annual result.

Python supplies only the source data (plan, actual and each department's
deviation over the months of actuals). Not one total is written in as a number:
anything that can be computed is computed by an Excel formula. That is what
makes the workbook a counterpart rather than a copy — it checks the dashboard
instead of repeating its output.

Correspondence with generate_dataset.py:
  plan_month[dep][m]   -> sheet "Данные", plan column
  fact                 -> sheet "Данные", actual column
  forecast             -> sheet "Данные", formula: plan * (1 + trend / 100)
  dev[dep]             -> sheet "Тренд"
  annual result        -> sheet "Итоги"

The plan is payroll (base pay plus variable pay) WITHOUT employer social
contributions: in Russian practice "ФОТ" excludes them.

IMPORTANT: openpyxl only writes formulas. What they compute is shown only by a
recalculation in real Excel — see recalc_model.ps1. Without that step, "it
agrees" is backed by nothing.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))

# --- Styling: the same colours as the first artefact ------------------------
H_FILL = PatternFill("solid", fgColor="1F3B57")
H_FONT = Font(color="FFFFFF", bold=True, size=10)
IN_FILL = PatternFill("solid", fgColor="FFF4D6")    # source data
CALC_FILL = PatternFill("solid", fgColor="EAF3FA")  # computed by formula
SET_FILL = PatternFill("solid", fgColor="E6F4EA")   # a setting, meant to be changed
TITLE = Font(bold=True, size=13)
BOLD = Font(bold=True)
MUTED = Font(color="6B7280", size=9)
THIN = Side(style="thin", color="D0D7DE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RUB = '#,##0.00'
PCT = '0.00%'

FACT_MONTHS = 8
ALERT_THRESHOLD = 0.03

# Sheet names. Formulas always quote them: Cyrillic names are safe either way,
# and a quoted reference survives a later rename to a name with a space.
S_SET, S_TREND, S_DATA, S_MON, S_DEP, S_SUM, S_LEG = (
    "Настройки", "Тренд", "Данные", "Месяцы", "Подразделения", "Итоги", "Пояснения")


def q(sheet):
    return "'%s'!" % sheet


def header(ws, row, labels, widths):
    for j, (label, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=j, value=label)
        c.fill = H_FILL
        c.font = H_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30


def note(ws, ref, text, merge):
    ws[ref] = text
    ws[ref].font = MUTED
    ws[ref].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(merge)


def read_dataset():
    path = os.path.join(HERE, "budget-plan-fact.csv")
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def read_deviation():
    # The deviations are read from the file the generator wrote rather than
    # derived back from the actuals: the actuals carry month-to-month noise, and
    # deriving backwards yields an average instead of the original figure.
    dev = {}
    with open(os.path.join(HERE, "department-deviation.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            dev[r["department"]] = float(r["run_rate_deviation_pct"])
    return dev


def build(out_path):
    rows = read_dataset()
    dev = read_deviation()

    departments = []
    months = []
    for r in rows:
        if r["department"] not in departments:
            departments.append(r["department"])
        if r["month"] not in months:
            months.append(r["month"])

    wb = Workbook()

    # ======================================================================
    # SHEET Настройки
    # ======================================================================
    st = wb.active
    st.title = S_SET
    st["A1"] = "Настройки модели"
    st["A1"].font = TITLE
    st["A2"] = ("План-факт ФОТ. Данные смоделированы, это иллюстрация метода. "
                "Comp Budget Lab.")
    st["A2"].font = MUTED
    st["A4"] = "Менять можно только зелёные ячейки. Всё остальное пересчитается само."
    st["A4"].font = MUTED

    settings = [
        ("fact_months", FACT_MONTHS, RUB.replace(".00", ""), "Месяцев с фактом",
         "Месяцы нумеруются с января: 8 значит, что факт есть по август "
         "включительно. Остальные месяцы — прогноз."),
        ("alert_threshold", ALERT_THRESHOLD, '0.0%', "Порог сигнала",
         "Если ожидание на год выше плана на этот процент или больше, на листе "
         "«Итоги» загорается «проблема»."),
    ]
    header(st, 6, ["Имя", "Значение", "Что это", "Пояснение"], [18, 12, 22, 80])
    for i, (name, val, fmt, label, text) in enumerate(settings):
        r = 7 + i
        st.cell(row=r, column=1, value=name).border = BOX
        c = st.cell(row=r, column=2, value=val)
        c.fill, c.border, c.font, c.number_format = SET_FILL, BOX, BOLD, fmt
        st.cell(row=r, column=3, value=label).border = BOX
        d = st.cell(row=r, column=4, value=text)
        d.border = BOX
        d.alignment = Alignment(wrap_text=True, vertical="top")
        wb.defined_names.add(DefinedName(name, attr_text="%s$B$%d" % (q(S_SET), r)))
        st.row_dimensions[r].height = 28

    st["A11"] = "Что в плане"
    st["A11"].font = BOLD
    note(st, "A12",
         "План — ФОТ подразделения: оклады и переменная часть, без страховых взносов "
         "работодателя. Годовая сумма разложена по месяцам с учётом сезонности "
         "отпусков. Прогноз на оставшиеся месяцы продолжает тренд каждого "
         "подразделения: кто восемь месяцев шёл выше плана, тот и дальше идёт выше.",
         "A12:D14")

    # ======================================================================
    # SHEET Тренд — each department's deviation over the months of actuals
    # ======================================================================
    tr = wb.create_sheet(S_TREND)
    tr["A1"] = "Тренд подразделений"
    tr["A1"].font = TITLE
    note(tr, "A2",
         "На сколько процентов подразделение шло выше (или ниже) плана в месяцы "
         "с фактом. Прогноз на оставшиеся месяцы — план, умноженный на (1 + тренд). "
         "Это исходные данные: поменяйте тренд, и прогноз с итогами пересчитаются.",
         "A2:C3")
    THDR = 5
    header(tr, THDR, ["Подразделение", "Тренд, %"], [24, 14])
    trend_row = {}
    for i, dep in enumerate(departments):
        r = THDR + 1 + i
        tr.cell(row=r, column=1, value=dep).border = BOX
        # Full precision: the forecast multiplies it across months and
        # departments, and a truncated decimal grows into kopecks on the year.
        c = tr.cell(row=r, column=2, value=dev[dep])
        c.fill, c.border, c.number_format = IN_FILL, BOX, "0.00"
        trend_row[dep] = r
    tr.freeze_panes = "A%d" % (THDR + 1)

    # ======================================================================
    # SHEET Данные — plan, actual, forecast, deviation, running total
    # ======================================================================
    ws = wb.create_sheet(S_DATA)
    ws["A1"] = "План и факт по подразделениям и месяцам"
    ws["A1"].font = TITLE
    ws["A2"] = ("Жёлтое — из CSV. Голубое — считается здесь. Отклонение — факт "
                "(или прогноз) минус план: плюс значит перерасход.")
    ws["A2"].font = MUTED

    HDR = 4
    cols = ["Подразделение", "Месяц", "№", "План, ₽", "Факт, ₽", "Прогноз, ₽",
            "Ожидание, ₽", "Отклонение, ₽", "Отклонение, %", "Накопительно, ₽"]
    header(ws, HDR, cols, [22, 8, 5, 15, 15, 15, 15, 15, 13, 16])

    notes = {
        4: "ФОТ без страховых взносов, разложенный по месяцам с учётом сезонности "
           "отпусков: выше в начале и конце года, ниже летом.",
        5: "Факт за месяцы 1..fact_months. Дальше пусто.",
        6: "Прогноз на оставшиеся месяцы: план, умноженный на (1 + тренд "
           "подразделения с листа «Тренд»).",
        7: "То из двух, что известно за месяц: факт, если он есть, иначе прогноз.",
        10: "Нарастающий итог отклонения внутри подразделения с января. Показывает, "
            "копятся месячные отклонения за год или гасят друг друга.",
    }
    for col, text in notes.items():
        ws.cell(row=HDR, column=col).comment = Comment(text, "Comp Budget Lab", height=120, width=320)

    by_key = {(r["department"], int(r["month_no"])): r for r in rows}
    first_data = HDR + 1
    r = first_data
    dep_first = {}
    for dep in departments:
        dep_first[dep] = r
        for m in range(1, 13):
            src = by_key[(dep, m)]
            ws.cell(row=r, column=1, value=dep).border = BOX
            ws.cell(row=r, column=2, value=src["month"]).border = BOX
            ws.cell(row=r, column=3, value=m).border = BOX

            c = ws.cell(row=r, column=4, value=float(src["plan"]))
            c.fill, c.border, c.number_format = IN_FILL, BOX, RUB

            if src["fact"]:
                c = ws.cell(row=r, column=5, value=float(src["fact"]))
                c.fill, c.border, c.number_format = IN_FILL, BOX, RUB
            else:
                ws.cell(row=r, column=5).border = BOX

            # ROUND to the kopeck sits where the generator's round(forecast, 2)
            # does, rather than being papered over by a wider tolerance.
            ws.cell(row=r, column=6,
                    value=("=IF(C{r}<=fact_months,\"\","
                           "ROUND(D{r}*(1+{t}$B${tr}/100),2))"
                           ).format(r=r, t=q(S_TREND), tr=trend_row[dep]))
            c = ws.cell(row=r, column=6)
            c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB

            c = ws.cell(row=r, column=7,
                        value="=IF(C{r}<=fact_months,E{r},F{r})".format(r=r))
            c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB

            c = ws.cell(row=r, column=8, value="=G{r}-D{r}".format(r=r))
            c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB

            c = ws.cell(row=r, column=9, value="=IF(D{r}=0,\"\",H{r}/D{r})".format(r=r))
            c.fill, c.border, c.number_format = CALC_FILL, BOX, PCT

            formula = "=H{r}".format(r=r) if m == 1 else "=J{p}+H{r}".format(p=r - 1, r=r)
            c = ws.cell(row=r, column=10, value=formula)
            c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB

            r += 1

    last_data = r - 1
    ws.freeze_panes = "A%d" % first_data
    ws.auto_filter.ref = "A%d:J%d" % (HDR, last_data)
    D = q(S_DATA)

    # ======================================================================
    # SHEET Месяцы — company totals by month
    # ======================================================================
    mo = wb.create_sheet(S_MON)
    mo["A1"] = "Итог по компании за каждый месяц"
    mo["A1"].font = TITLE
    note(mo, "A2",
         "Две разные величины. Месячное отклонение говорит, насколько этот месяц "
         "ушёл от плана; накопительное несёт в себе все месяцы до него. Месяц может "
         "стать лучше, а накопительное — продолжать расти.",
         "A2:G3")

    MHDR = 5
    header(mo, MHDR, ["Месяц", "№", "План, ₽", "Ожидание, ₽",
                      "Отклонение, ₽", "Отклонение, %", "Накопительно, ₽"],
           [10, 5, 16, 16, 16, 13, 16])

    for i, month_name in enumerate(months):
        r = MHDR + 1 + i
        m = i + 1
        mo.cell(row=r, column=1, value=month_name).border = BOX
        mo.cell(row=r, column=2, value=m).border = BOX
        for col, src_col in ((3, "D"), (4, "G")):
            c = mo.cell(row=r, column=col,
                        value="=SUMIF({d}$C${f}:$C${l},$B{r},{d}${sc}${f}:${sc}${l})"
                              .format(d=D, f=first_data, l=last_data, r=r, sc=src_col))
            c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = mo.cell(row=r, column=5, value="=D{r}-C{r}".format(r=r))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = mo.cell(row=r, column=6, value="=E{r}/C{r}".format(r=r))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, PCT
        formula = "=E{r}".format(r=r) if m == 1 else "=G{p}+E{r}".format(p=r - 1, r=r)
        c = mo.cell(row=r, column=7, value=formula)
        c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB

    m_first, m_last = MHDR + 1, MHDR + 12
    year_row = m_last + 1
    mo.cell(row=year_row, column=1, value="ГОД").font = BOLD
    for col in (3, 4, 5):
        L = get_column_letter(col)
        c = mo.cell(row=year_row, column=col,
                    value="=SUM(%s%d:%s%d)" % (L, m_first, L, m_last))
        c.font, c.border, c.number_format = BOLD, BOX, RUB
    c = mo.cell(row=year_row, column=6, value="=E{r}/C{r}".format(r=year_row))
    c.font, c.border, c.number_format = BOLD, BOX, PCT
    mo.freeze_panes = "A%d" % m_first
    M = q(S_MON)

    # ======================================================================
    # SHEET Подразделения — totals by department
    # ======================================================================
    dp = wb.create_sheet(S_DEP)
    dp["A1"] = "По подразделениям: месяцы с фактом и год"
    dp["A1"].font = TITLE
    dp["A2"] = ("Знак отклонения внутри подразделения устойчив: оно весь год либо "
                "выше плана, либо ниже. Перерасход обычно сидит в одном месте.")
    dp["A2"].font = MUTED

    DHDR = 4
    header(dp, DHDR, ["Подразделение", "План с начала года, ₽", "Факт с начала года, ₽",
                      "Отклонение, ₽", "Отклонение, %", "План на год, ₽",
                      "Ожидание на год, ₽", "Отклонение за год, ₽", "За год, %"],
           [22, 17, 17, 16, 13, 16, 17, 17, 10])

    for i, dep in enumerate(departments):
        r = DHDR + 1 + i
        f, l = dep_first[dep], dep_first[dep] + 11
        dp.cell(row=r, column=1, value=dep).border = BOX
        for col, src_col in ((2, "D"), (3, "E")):
            c = dp.cell(row=r, column=col,
                        value="=SUMIF({d}$C${f}:$C${l},\"<=\"&fact_months,{d}${sc}${f}:${sc}${l})"
                              .format(d=D, f=f, l=l, sc=src_col))
            c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = dp.cell(row=r, column=4, value="=C{r}-B{r}".format(r=r))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = dp.cell(row=r, column=5, value="=D{r}/B{r}".format(r=r))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, PCT
        c = dp.cell(row=r, column=6, value="=SUM({d}$D${f}:$D${l})".format(d=D, f=f, l=l))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = dp.cell(row=r, column=7, value="=SUM({d}$G${f}:$G${l})".format(d=D, f=f, l=l))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = dp.cell(row=r, column=8, value="=G{r}-F{r}".format(r=r))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, RUB
        c = dp.cell(row=r, column=9, value="=H{r}/F{r}".format(r=r))
        c.fill, c.border, c.number_format = CALC_FILL, BOX, PCT

    d_first, d_last = DHDR + 1, DHDR + len(departments)
    r = d_last + 1
    dp.cell(row=r, column=1, value="ИТОГО").font = BOLD
    for col in (2, 3, 4, 6, 7, 8):
        L = get_column_letter(col)
        c = dp.cell(row=r, column=col, value="=SUM(%s%d:%s%d)" % (L, d_first, L, d_last))
        c.font, c.border, c.number_format = BOLD, BOX, RUB
    for col, num in ((5, "=D{r}/B{r}"), (9, "=H{r}/F{r}")):
        c = dp.cell(row=r, column=col, value=num.format(r=r))
        c.font, c.border, c.number_format = BOLD, BOX, PCT
    dp.freeze_panes = "A%d" % d_first

    # ======================================================================
    # SHEET Итоги — what the dashboard shows on its first screen
    # ======================================================================
    sm = wb.create_sheet(S_SUM)
    sm["A1"] = "Итоги"
    sm["A1"].font = TITLE
    sm["A2"] = ("Каждая цифра здесь — формула. Поменяйте факт на листе «Данные» "
                "и посмотрите, как всё пересчитается.")
    sm["A2"].font = MUTED
    sm.column_dimensions["A"].width = 46
    for col in ("B", "C"):
        sm.column_dimensions[col].width = 20

    # Rows are declared by key, and references are assembled from the actual
    # addresses. Row numbers are never typed by hand: that is how a "Total"
    # once ended up referring to itself.
    SUM_ROWS = [
        ("h1",        "Месяцы с фактом", None, None),
        ("plan8",     "План с начала года",
         "=SUMIF({M}$B${mf}:$B${ml},\"<=\"&fact_months,{M}$C${mf}:$C${ml})", RUB),
        ("fact8",     "Факт с начала года",
         "=SUMIF({M}$B${mf}:$B${ml},\"<=\"&fact_months,{M}$D${mf}:$D${ml})", RUB),
        ("dev8",      "Отклонение", "={fact8}-{plan8}", RUB),
        ("dev8pct",   "В % к плану", "={dev8}/{plan8}", PCT),
        ("gap1",      "", None, None),
        ("h2",        "Оставшиеся месяцы", None, None),
        ("plan4",     "План", "={yplan}-{plan8}", RUB),
        ("fc4",       "Прогноз", "={yexp}-{fact8}", RUB),
        ("dev4",      "Отклонение", "={fc4}-{plan4}", RUB),
        ("dev4pct",   "В % к плану", "={dev4}/{plan4}", PCT),
        ("gap2",      "", None, None),
        ("h3",        "Год целиком", None, None),
        ("yplan",     "План", "={M}$C${yr}", RUB),
        ("yexp",      "Ожидание (факт + прогноз)", "={M}$D${yr}", RUB),
        ("ydelta",    "Отклонение", "={yexp}-{yplan}", RUB),
        ("ydeltapct", "В % к плану", "={ydelta}/{yplan}", PCT),
        ("signal",    "Сигнал",
         "=IF({ydeltapct}>=alert_threshold,\"проблема\",\"в пределах порога\")", None),
        ("gap3",      "", None, None),
        ("cross",     "Первый месяц, когда накопительно выше плана",
         "=IFERROR(INDEX({M}$A${mf}:$A${ml},MATCH(TRUE,INDEX({M}$G${mf}:$G${ml}>0,0),0)),\"нет\")",
         None),
    ]

    addr = {key: "B%d" % (4 + i) for i, (key, _, _, _) in enumerate(SUM_ROWS)}
    ctx = dict(addr, M=M, mf=m_first, ml=m_last, yr=year_row)

    for i, (key, label, formula, fmt) in enumerate(SUM_ROWS):
        r = 4 + i
        c = sm.cell(row=r, column=1, value=label)
        if formula is None and label:
            c.font = BOLD
        if formula:
            v = sm.cell(row=r, column=2, value=formula.format(**ctx))
            v.fill, v.border = CALC_FILL, BOX
            if fmt:
                v.number_format = fmt

    # The addresses recalc_model.ps1 reads. They are written here rather than
    # remembered in that script: otherwise inserting a row makes them drift
    # silently. D1 also marks this sheet for the script, which finds it by
    # this marker rather than by its Cyrillic name.
    sm["D1"] = "cells_for_check"
    sm["D1"].font = MUTED
    check_keys = ("plan8", "fact8", "fc4", "yplan", "yexp", "ydelta", "cross", "signal")
    for j, key in enumerate(check_keys):
        sm.cell(row=2 + j, column=4, value=key).font = MUTED
        sm.cell(row=2 + j, column=5, value=addr[key]).font = MUTED
    sm.column_dimensions["D"].width = 16
    sm.column_dimensions["E"].width = 8

    last_sum = 4 + len(SUM_ROWS)
    note(sm, "A%d" % (last_sum + 1),
         "Отклонение везде — ожидание минус план. Плюс значит, что компания потратила "
         "больше плана, минус — меньше. Столбцы D и E — служебные: по ним проверочный "
         "скрипт находит итоговые ячейки.",
         "A%d:C%d" % (last_sum + 1, last_sum + 3))

    # ======================================================================
    # SHEET Пояснения
    # ======================================================================
    lg = wb.create_sheet(S_LEG)
    lg["A1"] = "Пояснения и метод"
    lg["A1"].font = TITLE
    lg.column_dimensions["A"].width = 32
    lg.column_dimensions["B"].width = 110

    blocks = [
        ("Цвета", ""),
        ("Жёлтый", "Исходные данные. Берутся из CSV, не считаются."),
        ("Голубой", "Считается формулой. Эти ячейки пересчитываются сами."),
        ("Зелёный", "Настройка на листе «Настройки». Их можно менять."),
        ("", ""),
        ("Порядок расчёта", ""),
        ("1. План",
         "ФОТ подразделения на год — оклады и переменная часть, без страховых взносов "
         "работодателя. Разложен на двенадцать месяцев с учётом сезонности отпусков: "
         "выше в начале и конце года, ниже летом. Коэффициенты сезонности в сумме дают "
         "двенадцать, поэтому сезонность двигает деньги внутри года, не меняя годовую сумму."),
        ("2. Факт",
         "Месяцы 1..fact_months. Знак отклонения в подразделении устойчив: оно весь год "
         "либо выше плана, либо ниже. Одно подразделение заметно выбивается из ряда: "
         "бюджет редко расходится ровно, он расходится в одном месте."),
        ("3. Прогноз",
         "На оставшиеся месяцы — план, умноженный на (1 + тренд подразделения). Тренд — "
         "отклонение подразделения от плана в месяцы с фактом, лист «Тренд». Прогноз "
         "«ровно по плану» молча предполагал бы, что перерасход с сентября исчезнет."),
        ("4. Отклонение",
         "Ожидание минус план. Плюс — перерасход. Показано и по месяцам, и нарастающим "
         "итогом с января."),
        ("5. Год и сигнал",
         "План против факта и прогноза. Если ожидание на год выше плана на порог или "
         "больше (настройка alert_threshold, по умолчанию 3%), сигнал — «проблема»."),
        ("", ""),
        ("Две величины, которые не надо путать", ""),
        ("Месячное отклонение",
         "Насколько один месяц ушёл от плана. Отвечает на вопрос: этот месяц хуже прошлого?"),
        ("Накопительное отклонение",
         "То же самое, сложенное с января. Отвечает на другой вопрос: какая часть годового "
         "бюджета уже ушла сверх плана. Месяц может улучшиться, а накопительное — расти. "
         "Обе цифры верны; одну нельзя читать вместо другой."),
        ("", ""),
        ("Чего эта модель не делает", ""),
        ("Не объясняет отклонения",
         "Показывает, где и сколько, но не почему. Причины выясняются в разговоре "
         "с руководителем подразделения, а не на экране."),
        ("Страховые взносы",
         "В план и факт не входят: ФОТ здесь — оклады и переменная часть. Если ваш "
         "бюджет ведётся вместе со взносами, загрузите суммы со взносами — модель "
         "считает отклонения от того, что в файле."),
        ("Данные смоделированы",
         "Получены генератором со сидом, а не из реальной компании. Модель показывает "
         "подход к контролю бюджета и не является юридической или бухгалтерской консультацией."),
    ]
    r = 3
    for label, text in blocks:
        if label:
            c = lg.cell(row=r, column=1, value=label)
            c.font = BOLD if not text else Font(bold=False)
            c.alignment = Alignment(vertical="top")
        if text:
            c = lg.cell(row=r, column=2, value=text)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            lg.row_dimensions[r].height = max(14, 13 * (len(text) // 95 + 1))
        r += 1

    wb.save(out_path)
    return out_path, len(rows), departments


def main():
    # The workbook sits beside the page, not in build/: it is one of the files
    # a reader takes away.
    out = os.path.join(os.path.dirname(HERE), "budget-plan-fact-model.xlsx")
    path, n, departments = build(out)
    print("Written: %s" % path)
    print("Data rows: %d, departments: %d" % (n, len(departments)))
    print()
    print("NOTE: openpyxl has only written the formulas. Nothing has been")
    print("computed yet. The next step is required:")
    print("  powershell -File recalc_model.ps1")


if __name__ == "__main__":
    main()
